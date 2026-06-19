﻿"""文件处理模块 - 支持 PDF、Word、TXT 等文件的解析与内容提取"""

import uuid
from datetime import datetime
from io import BytesIO
from pathlib import Path
from typing import Optional

from src.config import Settings, logger

UPLOAD_DIR = Settings.PROJECT_ROOT / "data" / "uploads"
UPLOAD_DIR.mkdir(parents=True, exist_ok=True)

MAX_FILE_SIZE = 10 * 1024 * 1024  # 10MB
ALLOWED_EXTENSIONS = {
    ".pdf", ".docx", ".doc", ".txt", ".md",
    ".py", ".java", ".cpp", ".c", ".h", ".js", ".ts",
    ".html", ".css", ".json", ".xml", ".yaml", ".yml", ".csv",
}

def allowed_file(filename: str) -> bool:
    ext = Path(filename).suffix.lower()
    return ext in ALLOWED_EXTENSIONS

def get_file_size_limit() -> int:
    return MAX_FILE_SIZE

def parse_pdf(file_bytes: bytes) -> str:
    try:
        from PyPDF2 import PdfReader
        reader = PdfReader(BytesIO(file_bytes))
        text_parts = []
        for i, page in enumerate(reader.pages):
            page_text = page.extract_text()
            if page_text:
                text_parts.append(f"[Page {i+1}]\n{page_text.strip()}")
        result = "\n\n".join(text_parts)
        if not result.strip():
            raise ValueError("PDF file has no extractable text")
        logger.info(f"PDF parsed | pages={len(reader.pages)} chars={len(result)}")
        return result
    except ImportError:
        raise RuntimeError("PyPDF2 not installed")

def parse_docx(file_bytes: bytes) -> str:
    try:
        from docx import Document
        doc = Document(BytesIO(file_bytes))
        text_parts = []
        for para in doc.paragraphs:
            if para.text.strip():
                text_parts.append(para.text.strip())
        result = "\n\n".join(text_parts)
        if not result.strip():
            raise ValueError("Word file has no text content")
        logger.info(f"Word parsed | paras={len(doc.paragraphs)} chars={len(result)}")
        return result
    except ImportError:
        raise RuntimeError("python-docx not installed")

def parse_txt(file_bytes: bytes) -> str:
    try:
        text = file_bytes.decode("utf-8")
    except UnicodeDecodeError:
        try:
            text = file_bytes.decode("gbk")
        except UnicodeDecodeError:
            text = file_bytes.decode("latin-1", errors="replace")
    if not text.strip():
        raise ValueError("Text file is empty")
    logger.info(f"TXT parsed | chars={len(text)}")
    return text

def parse_code_file(file_bytes: bytes, filename: str) -> str:
    ext = Path(filename).suffix.lower()
    lang_map = {
        ".py": "python", ".java": "java", ".cpp": "cpp", ".c": "c", ".h": "c",
        ".js": "javascript", ".ts": "typescript", ".html": "html", ".css": "css",
        ".json": "json", ".xml": "xml", ".yaml": "yaml", ".yml": "yaml",
        ".csv": "csv", ".md": "markdown",
    }
    lang = lang_map.get(ext, "")
    text = parse_txt(file_bytes)
    return f"```{lang}\n{text}\n```"

def parse_file(file_bytes: bytes, filename: str) -> str:
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        return parse_pdf(file_bytes)
    elif ext in (".docx", ".doc"):
        return parse_docx(file_bytes)
    elif ext == ".txt":
        return parse_txt(file_bytes)
    elif ext in (".py", ".java", ".cpp", ".c", ".h", ".js", ".ts", ".html",
                 ".css", ".json", ".xml", ".yaml", ".yml", ".csv", ".md"):
        return parse_code_file(file_bytes, filename)
    else:
        raise ValueError(f"Unsupported format: {ext}")

def save_uploaded_file(file_bytes: bytes, filename: str) -> dict:
    file_id = uuid.uuid4().hex[:12]
    safe_name = f"{file_id}_{Path(filename).name}"
    file_path = UPLOAD_DIR / safe_name
    with open(file_path, "wb") as f:
        f.write(file_bytes)
    return {
        "file_id": file_id, "original_name": filename,
        "saved_name": safe_name, "path": str(file_path),
        "size": len(file_bytes), "upload_time": datetime.now().isoformat(),
    }

def get_uploaded_file(file_id: str) -> Optional[dict]:
    for f in UPLOAD_DIR.iterdir():
        if f.is_file() and f.name.startswith(file_id + "_"):
            return {
                "file_id": file_id,
                "original_name": f.name[len(file_id) + 1:],
                "saved_name": f.name, "path": str(f),
                "size": f.stat().st_size,
            }
    return None

def delete_uploaded_file(file_id: str) -> bool:
    for f in UPLOAD_DIR.iterdir():
        if f.is_file() and f.name.startswith(file_id + "_"):
            f.unlink()
            return True
    return False
